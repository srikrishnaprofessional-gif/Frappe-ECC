#!/usr/bin/env python3
"""
Generates the comprehensive FRAPPE_ECC_SOP.docx and converts it to FRAPPE_ECC_SOP.pdf.
"""

import os
import sys
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = Path(__file__).resolve().parent.parent / "docs"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE", alert_type="note"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=120, bottom=120, left=200, right=200)
    
    border_colors = {
        "note": "2B579A",
        "warning": "D83B01",
        "tip": "107C41"
    }
    bg_colors = {
        "note": "F0F4F8",
        "warning": "FFF4CE",
        "tip": "EDF7ED"
    }
    
    fill_hex = bg_colors.get(alert_type, "F0F4F8")
    b_color = border_colors.get(alert_type, "2B579A")
    set_cell_background(cell, fill_hex)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{b_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor.from_string(b_color)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(code_text.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(36, 41, 46)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx(output_path):
    doc = docx.Document()
    
    # Page setup - Standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "STANDARD OPERATING PROCEDURE | FRAPPE ECC (ENGINEERING COORDINATION CENTER)"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Calibri"
        hp.style.font.size = Pt(8.5)
        hp.style.font.color.rgb = RGBColor(128, 128, 128)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Document ID: SOP-ENG-FRAPPE-ECC-001  |  Confidential & Internal Engineering Standard"
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.style.font.name = "Calibri"
        fp.style.font.size = Pt(8.5)
        fp.style.font.color.rgb = RGBColor(128, 128, 128)

    # 1. Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("STANDARD OPERATING PROCEDURE (SOP)")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.color.rgb = RGBColor(31, 78, 121) # Professional navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Engineering Coordination Center for Frappe Framework & ERPNext (Frappe ECC)")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(89, 89, 89)

    # 2. Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document ID", "SOP-ENG-FRAPPE-ECC-001"),
        ("Version & Release", "Version 1.0.0 (Production Release)"),
        ("Effective Date", "2026-09-18"),
        ("Applicability", "Frappe v14 / v15 / v16, ERPNext, Antigravity, Claude Code, Cursor")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "E9EEF4")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_sop_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        return h

    # Section 1
    add_sop_heading("1. Purpose & Scope", level=1)
    p = doc.add_paragraph(
        "This Standard Operating Procedure defines the mandatory engineering practices, execution phases, "
        "and automated quality controls for building, testing, and securing Frappe Framework applications using "
        "Frappe ECC. Frappe ECC turns your AI coding assistant (Antigravity, Claude Code, Cursor, Codex) into a "
        "fully coordinated engineering system equipped with 12 specialized agents, 14 production workflow skills, "
        "13 slash commands, and Frappe Shield."
    )
    p.paragraph_format.space_after = Pt(6)

    # Section 2
    add_sop_heading("2. System Architecture & Component Inventory", level=1)
    
    comp_table = doc.add_table(rows=6, cols=3)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Component Layer", "Count", "Core Responsibilities"]
    for col_idx, h in enumerate(headers):
        cell = comp_table.rows[0].cells[col_idx]
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 80, 80, 100, 100)
        cp = cell.paragraphs[0]
        cr = cp.add_run(h)
        cr.bold = True
        cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(255, 255, 255)
        cr.font.size = Pt(10)

    rows_data = [
        ("Specialist AI Agents", "12", "Planner, Architect, Controller Builder, Desk UI Builder, TDD Guide, Code Reviewer, Security Reviewer, Bench DevOps, Migration Patcher, Report Builder, API Integrator, Doc Updater"),
        ("Workflow Skills", "14", "DocType modeling, QueryBuilder/ORM, hooks.py, Desk Client Scripts, REST API, Permissions, TDD with FrappeTestCase, RQ background jobs, Script Reports, Patches, Bench CLI, Frappe UI, Web Portal, Security Audit"),
        ("Slash Commands", "13", "Quick entry points: /frappe:plan, /frappe:doctype, /frappe:controller, /frappe:client-script, /frappe:hook, /frappe:api, /frappe:test, /frappe:review, /frappe:security, /frappe:patch, /frappe:report, /frappe:bench, /frappe:help"),
        ("Coding Rules", "5", "Always-loaded standards: Core architecture, Python backend, Desk JS, Security/Permissions, and MariaDB/PostgreSQL database optimization"),
        ("Frappe Shield", "1 CLI", "AST-based static analyzer detecting SQL injection in frappe.db.sql, db.commit() transaction violations, unvalidated guest endpoints, and N+1 query bottlenecks")
    ]
    for row_idx, r_data in enumerate(rows_data, 1):
        row = comp_table.rows[row_idx]
        bg = "F9FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate(r_data):
            c = row.cells[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 100, 100)
            cp = c.paragraphs[0]
            cr = cp.add_run(text)
            cr.font.name = "Calibri"
            cr.font.size = Pt(9.5)
            if c_idx == 0:
                cr.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 3
    add_sop_heading("3. Installation & Verification Procedure", level=1)
    doc.add_paragraph("Follow these exact commands to install Frappe ECC into your code editor:")
    
    add_callout(doc, "Antigravity installation automatically registers skills into your global config directory so the AI assistant recognizes all Frappe paradigms across every project.", "RECOMMENDED SETUP", "tip")
    
    doc.add_paragraph("Command for Linux / macOS / WSL:")
    add_code_block(doc, "./install.sh --profile minimal --target antigravity")
    
    doc.add_paragraph("Command for Windows PowerShell:")
    add_code_block(doc, ".\\install.ps1 -Profile minimal -Target antigravity")
    
    doc.add_paragraph("Verification Check:")
    add_code_block(doc, "python bin/frappe_ecc_install.py doctor")

    # Section 4
    add_sop_heading("4. The 10-Phase Feature Development SOP", level=1)
    doc.add_paragraph("Every engineer and AI assistant MUST adhere to the following 10-phase sequence:")

    phases = [
        ("Phase 1: Feature Planning & Architecture Blueprinting",
         "Never write code without a blueprint. Invoke the planner agent using '/frappe:plan' or ask your assistant. "
         "The planner determines DocType taxonomy (Standard vs Child Table vs Single vs Submittable), fieldtypes, "
         "autoname rules (format:AST-.YYYY.-.#####), module boundaries, and execution order."),

        ("Phase 2: Schema Modeling & DocType Scaffolding",
         "Generate schema JSON, controller, client script, and test files using '/frappe:doctype'. "
         "Ensure search indexes ('search_index': 1) are applied to frequent lookup fields. "
         "Embed child tables via 'fieldtype': 'Table' and 'options': '<Child DocType>'."),

        ("Phase 3: Test-Driven Development (TDD) — Write Tests First",
         "Before implementing business logic, write failing unit tests in 'test_<doctype>.py' with FrappeTestCase. "
         "Assert autoname prefix, mandatory fields, validation failures, and transaction rollbacks. "
         "Run tests using 'bench run-tests --doctype <DocType>' to confirm RED failing state."),

        ("Phase 4: Backend Controller & Lifecycle Implementation",
         "Implement controller logic using '/frappe:controller'. Follow Frappe's exact lifecycle: before_insert, "
         "validate, on_submit, and on_cancel. CRITICAL: NEVER call frappe.db.commit() inside controller events. "
         "Use frappe.qb (QueryBuilder) for complex queries. Wrap user strings with _('...')."),

        ("Phase 5: Desk UI & Reactive Client Scripting",
         "Construct responsive Desk forms using '/frappe:client-script'. Use 'frappe.ui.form.on' for events: "
         "onload, refresh, and field triggers. Use frm.add_custom_button, frm.set_df_property, and native dialogs. "
         "CRITICAL: Never manipulate the DOM directly using jQuery or raw selectors."),

        ("Phase 6: Hook Registration & Event Orchestration",
         "Connect application events, cron jobs, and class overrides into hooks.py using '/frappe:hook'. "
         "Configure doc_events for cross-doctype triggers, scheduler_events for cron routines, and fixtures "
         "for persistent custom fields and roles."),

        ("Phase 7: Whitelisted APIs & Webhook Integrations",
         "Create secure REST endpoints using '/frappe:api'. Annotate endpoints with @frappe.whitelist(). "
         "Enforce frappe.only_for() role authorization. Add @frappe.rate_limit on public endpoints. "
         "Sanitize inputs with frappe.utils.escape_html()."),

        ("Phase 8: Database Schema Migrations & Patches",
         "Create idempotent migration patches in 'patches/vX_Y/<patch>.py' using '/frappe:patch'. "
         "Always call frappe.reload_doc() before database writes. Verify with 'bench migrate'."),

        ("Phase 9: Fresh-Context Code Review",
         "Execute an independent, fresh-context code review using '/frappe:review'. "
         "Detects stray db.commit() calls, N+1 query bottlenecks, unindexed filters, and untranslated strings."),

        ("Phase 10: Security Audit with Frappe Shield",
         "Run automated AST security scanning using 'python bin/frappe-shield.py <path>' or '/frappe:security'. "
         "Flags SQL injection in frappe.db.sql, transaction violations, and direct docstatus mutations. "
         "All CRITICAL and HIGH issues must be resolved before deployment.")
    ]

    for p_title, p_desc in phases:
        add_sop_heading(p_title, level=2)
        doc.add_paragraph(p_desc).paragraph_format.space_after = Pt(4)

    # Section 5
    add_sop_heading("5. Frappe Shield Security Rules & Remediation", level=1)
    
    shield_table = doc.add_table(rows=6, cols=3)
    shield_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Vulnerability Pattern", "Severity", "Remediation Standard"]
    for c_idx, h in enumerate(s_headers):
        cell = shield_table.rows[0].cells[c_idx]
        set_cell_background(cell, "800000") # Dark red header
        set_cell_margins(cell, 80, 80, 100, 100)
        cp = cell.paragraphs[0]
        cr = cp.add_run(h)
        cr.bold = True
        cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(255, 255, 255)
        cr.font.size = Pt(10)

    s_rows = [
        ("SQL Injection: f-string or % in frappe.db.sql()", "CRITICAL", "Replace string interpolation with values={'key': val} or migrate to frappe.qb."),
        ("Manual DB Commit Inside Controller Hook", "CRITICAL", "Remove frappe.db.commit(). Frappe manages transaction commits automatically."),
        ("Unrestricted Guest Whitelist API", "HIGH", "Add @frappe.rate_limit(limit=10, seconds=60) and validate/sanitize all inputs."),
        ("Direct docstatus = 1 Assignment", "HIGH", "Use doc.submit() or doc.cancel() instead of directly overwriting docstatus."),
        ("N+1 Query Loop (frappe.get_doc in loop)", "MEDIUM", "Batch fetch using frappe.get_all(filters={'name': ['in', ids]}) or frappe.db.get_values().")
    ]
    for r_idx, r_data in enumerate(s_rows, 1):
        row = shield_table.rows[r_idx]
        bg = "FFF5F5" if "CRITICAL" in r_data[1] else "FFFFFF"
        for c_idx, text in enumerate(r_data):
            c = row.cells[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 100, 100)
            cp = c.paragraphs[0]
            cr = cp.add_run(text)
            cr.font.name = "Calibri"
            cr.font.size = Pt(9.5)
            if c_idx == 1:
                cr.bold = True
                if text == "CRITICAL":
                    cr.font.color.rgb = RGBColor(192, 0, 0)
                elif text == "HIGH":
                    cr.font.color.rgb = RGBColor(230, 126, 34)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 6
    add_sop_heading("6. DevOps & Bench Troubleshooting Runbook", level=1)
    
    devops_items = [
        ("Redis Connection Failure (redis.exceptions.ConnectionError)",
         "Ensure Redis is running: 'sudo systemctl restart redis-server'. Verify socket path in config/redis_cache.conf."),
        ("Duplicate Entry Error on Schema Migration",
         "Inspect table using 'bench mariadb'. Remove or update duplicate rows prior to applying unique index constraint."),
        ("UI Changes Not Reflecting in Desk Form",
         "Clear Redis cache and rebuild assets: 'bench --site <site> clear-cache && bench build --app <app>'. Perform hard browser refresh (Ctrl+F5)."),
        ("Background Workers Stalled / Queues Backlogged",
         "Inspect queue health with 'bench doctor'. Restart RQ workers: 'bench worker --queue default,short,long'.")
    ]
    for issue, fix in devops_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r_issue = p.add_run(f"• {issue}: ")
        r_issue.bold = True
        r_issue.font.name = "Calibri"
        r_fix = p.add_run(fix)
        r_fix.font.name = "Calibri"

    # Section 7
    add_sop_heading("7. Quick Reference Command Cheat Sheet", level=1)
    
    cmd_table = doc.add_table(rows=14, cols=3)
    cmd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cmd_headers = ["Slash Command", "Specialist Agent", "Standard Operation"]
    for c_idx, h in enumerate(cmd_headers):
        cell = cmd_table.rows[0].cells[c_idx]
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 80, 80, 100, 100)
        cp = cell.paragraphs[0]
        cr = cp.add_run(h)
        cr.bold = True
        cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(255, 255, 255)
        cr.font.size = Pt(10)

    commands_summary = [
        ("/frappe:plan '<feature>'", "frappe-planner", "Architect feature schema blueprint and relationships"),
        ("/frappe:doctype '<Name>'", "frappe-planner", "Scaffold complete .json, .py, .js, and test files"),
        ("/frappe:controller '<Name>'", "frappe-backend-builder", "Implement controller lifecycle hooks and validations"),
        ("/frappe:client-script '<Name>'", "frappe-desk-builder", "Construct reactive Desk form scripts and custom dialogs"),
        ("/frappe:hook [event|cron]", "frappe-architect", "Wire doc_events, scheduler cron, or class overrides"),
        ("/frappe:api '<name>'", "frappe-api-integrator", "Generate secure whitelisted REST API endpoint"),
        ("/frappe:test '<Name>'", "frappe-tdd-guide", "Scaffold unit tests with FrappeTestCase"),
        ("/frappe:review [path]", "frappe-code-reviewer", "Review diffs in fresh context for Frappe anti-patterns"),
        ("/frappe:security [path]", "frappe-security-reviewer", "Execute AST vulnerability audit (SQLi, IDOR, XSS)"),
        ("/frappe:patch '<desc>'", "frappe-migration-patcher", "Draft idempotent database migration patch"),
        ("/frappe:report '<name>'", "frappe-report-builder", "Scaffold Script Report (Python backend + JS frontend)"),
        ("/frappe:bench [task]", "frappe-bench-devops", "Diagnose bench errors and site operations"),
        ("/frappe:help", "All Agents", "Display command cheat sheet and quick guidance")
    ]
    for r_idx, (cmd, agent, desc) in enumerate(commands_summary, 1):
        row = cmd_table.rows[r_idx]
        bg = "F9FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([cmd, agent, desc]):
            c = row.cells[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 50, 50, 80, 80)
            cp = c.paragraphs[0]
            cr = cp.add_run(text)
            cr.font.name = "Calibri"
            cr.font.size = Pt(9)
            if c_idx == 0:
                cr.font.name = "Consolas"
                cr.bold = True
                cr.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    
    # Sign-off box
    sign_table = doc.add_table(rows=1, cols=1)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_cell = sign_table.cell(0, 0)
    set_cell_background(s_cell, "EAECEE")
    set_cell_margins(s_cell, 100, 100, 150, 150)
    sp = s_cell.paragraphs[0]
    sr = sp.add_run("APPROVED & ENFORCED BY: Antigravity Systems Engineering Group\nDISTRIBUTION: Open Source / All Engineering Teams")
    sr.font.name = "Calibri"
    sr.font.size = Pt(9)
    sr.bold = True
    sr.font.color.rgb = RGBColor(80, 80, 80)

    doc.save(str(output_path))
    print(f"[SUCCESS] DOCX generated: {output_path}")

def convert_to_pdf(docx_path, pdf_path):
    ps_script = f"""
    $docx = "{str(docx_path).replace('\\', '\\\\')}"
    $pdf = "{str(pdf_path).replace('\\', '\\\\')}"
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    try {{
        $doc = $word.Documents.Open($docx)
        $doc.SaveAs([ref]$pdf, [ref]17)
        $doc.Close()
        Write-Output "SUCCESS"
    }} finally {{
        $word.Quit()
        [System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null
    }}
    """
    
    ps_file = DOCS_DIR / "_convert_tmp.ps1"
    with open(ps_file, "w", encoding="utf-8") as f:
        f.write(ps_script)
        
    try:
        res = subprocess.run(
            ["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ps_file)],
            capture_output=True,
            text=True,
            check=True
        )
        print(f"[SUCCESS] PDF generated: {pdf_path}")
    finally:
        if ps_file.exists():
            ps_file.unlink()

def main():
    docx_path = DOCS_DIR / "FRAPPE_ECC_SOP.docx"
    pdf_path = DOCS_DIR / "FRAPPE_ECC_SOP.pdf"
    
    build_docx(docx_path)
    convert_to_pdf(docx_path, pdf_path)
    
    # Also copy to user's Downloads directory for convenient access!
    user_downloads = Path.home() / "Downloads"
    if user_downloads.exists():
        import shutil
        shutil.copy2(docx_path, user_downloads / "FRAPPE_ECC_SOP.docx")
        shutil.copy2(pdf_path, user_downloads / "FRAPPE_ECC_SOP.pdf")
        print(f"[SUCCESS] Copied copies to Downloads: {user_downloads / 'FRAPPE_ECC_SOP.pdf'}")

if __name__ == "__main__":
    main()
