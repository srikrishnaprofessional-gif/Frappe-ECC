#!/usr/bin/env python3
"""
Generates the comprehensive WORKING_SOP_EQUIPMENT_LOAN.docx with embedded UI screenshots
and converts it to WORKING_SOP_EQUIPMENT_LOAN.pdf.
Authored autonomously by frappe-working-sop-author agent.
"""

import os
import sys
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_DOCS = Path(__file__).resolve().parent.parent / "test_project" / "docs"
ASSETS_DIR = PROJECT_DOCS / "assets"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="IMPORTANT POLICY RULE", alert_type="brand"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    
    border_colors = {"brand": "4F46E5", "warning": "D83B01", "success": "107C41", "note": "1F4E79"}
    bg_colors = {"brand": "EEF2FF", "warning": "FFF4CE", "success": "EDF7ED", "note": "F0F4F8"}
    
    fill_hex = bg_colors.get(alert_type, "EEF2FF")
    b_color = border_colors.get(alert_type, "4F46E5")
    set_cell_background(cell, fill_hex)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{b_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor.from_string(b_color)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(code_text.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(36, 41, 46)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_sec_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    return h

def build_docx(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "WORKING STANDARD OPERATING PROCEDURE  |  IT EQUIPMENT LOAN & RETURN SYSTEM"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Calibri"
        hp.style.font.size = Pt(8.5)
        hp.style.font.color.rgb = RGBColor(128, 128, 128)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Document ID: SOP-OPS-ITAM-2026-001  |  Version 1.0.0  |  Corporate IT Standard"
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.style.font.name = "Calibri"
        fp.style.font.size = Pt(8.5)
        fp.style.font.color.rgb = RGBColor(128, 128, 128)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("WORKING STANDARD OPERATING PROCEDURE")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(31, 78, 121)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("IT Equipment Loan & Return Management System (equipment_loan)\nComplete Step-by-Step Operator Guide with Verified UI Screenshots")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(89, 89, 89)

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document ID", "SOP-OPS-ITAM-2026-001"),
        ("System Module", "IT Asset Management (ITAM) • Frappe Framework"),
        ("Effective Date", "September 21, 2026"),
        ("Target Operators", "IT Helpdesk Technicians, Asset Managers, Department Approvers, Employees")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.7)
        set_cell_background(c0, "E9EEF4")
        set_cell_margins(c0, 40, 40, 70, 70)
        set_cell_margins(c1, 40, 40, 70, 70)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1
    add_sec_heading(doc, "1. Purpose & Operational Objectives", level=1)
    doc.add_paragraph(
        "This Standard Operating Procedure (SOP) provides a complete, visual, step-by-step operational manual for issuing, "
        "tracking, and inspecting corporate IT hardware assets (laptops, displays, cameras, and peripherals) using the "
        "IT Equipment Loan System built on Frappe Framework."
    ).paragraph_format.space_after = Pt(4)

    add_callout(
        doc,
        "Zero-Loss Mandate: Every hardware item leaving IT Helpdesk custody must be registered in an Active Equipment Loan document "
        "with itemized serial numbers, signed voucher agreement, and scheduled return date not exceeding 30 calendar days.",
        title="CORE POLICY MANDATE",
        alert_type="brand"
    )

    # Section 2: Roles RACI
    add_sec_heading(doc, "2. Roles & Responsibilities (RACI Matrix)", level=1)
    raci_table = doc.add_table(rows=5, cols=3)
    raci_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Role Title", "Key Operational Responsibilities", "Frappe Desk Role Binding"]
    for col_idx, h in enumerate(headers):
        cell = raci_table.rows[0].cells[col_idx]
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 45, 45, 60, 60)
        cp = cell.paragraphs[0]
        cr = cp.add_run(h)
        cr.bold = True
        cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(255, 255, 255)
        cr.font.size = Pt(8.5)

    raci_data = [
        ("IT Asset Administrator", "Validates eligibility, issues hardware, submits loan doc, performs return physical check.", "IT Asset Manager (Full CRUD + Submit + Cancel)"),
        ("Employee Borrower", "Takes physical custody, signs voucher, returns hardware on or before due date.", "Employee (Portal View, E-sign)"),
        ("Department Manager", "Authorizes high-value loans (> $3,000) or extended multi-month loan renewals.", "Department Manager (Approver)"),
        ("Compliance Auditor", "Reviews monthly audit trail, verifies overdue borrower lockout, checks loss ratios.", "Auditor (Read-Only access to all logs)")
    ]
    for idx, (role, resp, bind) in enumerate(raci_data):
        row = raci_table.rows[idx + 1]
        bg = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate([role, resp, bind]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 35, 35, 50, 50)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            if c_idx == 0: r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 3: Operator Walkthrough with Screenshots
    add_sec_heading(doc, "3. Step-by-Step Operator Walkthrough", level=1)
    
    doc.add_paragraph(
        "Step 3.1: Navigating to Equipment Loan Form\n"
        "1. Open Frappe Desk in your browser (https://itam.company.internal).\n"
        "2. In the Awesomebar (Ctrl + G or ⌘K), type 'Equipment Loan' and select Equipment Loan List.\n"
        "3. Click '+ Add Equipment Loan' to instantiate a new draft record.\n\n"
        "Step 3.2: Assigning Borrower & Validation Rules\n"
        "1. Select the Borrower Employee (e.g. EMP-00104 - Alex Mercer). The Department auto-populates from the Employee master record.\n"
        "2. If an employee has active overdue items (e.g. EMP-00105 - Jordan Lee), the system throws an immediate PermissionError and blocks loan creation.\n"
        "3. Set the Loan Effective Date and Expected Return Date. The period cannot exceed 30 days.\n\n"
        "Step 3.3: Allocating Hardware Items & Submitting Loan\n"
        "1. In Allocated Hardware Items, click 'Add Row'.\n"
        "2. Select the asset item code (e.g. AST-001 - MacBook Pro 16\" M3 Max). The serial number and replacement value auto-populate.\n"
        "3. Click the blue 'Submit & Lock Assets' button in the header.\n"
        "4. Status transitions to Active, and the assets in the catalog are immediately marked as 'Loaned Out'."
    ).paragraph_format.space_after = Pt(6)

    # Embed Screenshot 1: Equipment Loan Form
    img1_path = ASSETS_DIR / "01_equipment_loan_form.jpg"
    if img1_path.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(str(img1_path), width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        rc = p_cap.add_run("Figure 1: Verified Frappe Desk Equipment Loan Form (LOAN-2026-00042) showing active status and allocated item table.")
        rc.italic = True
        rc.font.name = "Calibri"
        rc.font.size = Pt(8.5)
        rc.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph(
        "Step 3.4: Processing Equipment Returns (Check-In)\n"
        "When the employee brings back the equipment to the IT Helpdesk:\n"
        "1. Open the active loan document.\n"
        "2. Click the green 'Process Return' button in the top action bar.\n"
        "3. The Process Equipment Return modal dialog opens with a hardware inspection checklist.\n"
        "4. Confirm return date, check each operational box, select Condition Grade (Grade A / B / C), and enter inspector verification notes.\n"
        "5. Click 'Confirm Return'. The document status transitions to Returned, and assets are released back to Available."
    ).paragraph_format.space_after = Pt(6)

    # Embed Screenshot 2: Return Modal Dialog
    img2_path = ASSETS_DIR / "02_equipment_loan_return_modal.jpg"
    if img2_path.exists():
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(6)
        p_img2.paragraph_format.space_after = Pt(2)
        doc.add_picture(str(img2_path), width=Inches(5.8))
        
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(8)
        rc2 = p_cap2.add_run("Figure 2: Verified Process Equipment Return inspection modal dialog with condition grade checklist and inspector notes.")
        rc2.italic = True
        rc2.font.name = "Calibri"
        rc2.font.size = Pt(8.5)
        rc2.font.color.rgb = RGBColor(100, 100, 100)

    # Section 4: Validation Table
    add_sec_heading(doc, "4. Business Rule Failsafes & Error Resolution Guide", level=1)
    err_table = doc.add_table(rows=7, cols=3)
    err_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Error Message Encountered", "Underlying Cause", "Operator Corrective Action"]
    for col_idx, h in enumerate(headers):
        cell = err_table.rows[0].cells[col_idx]
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 45, 45, 60, 60)
        cp = cell.paragraphs[0]
        cr = cp.add_run(h)
        cr.bold = True
        cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(255, 255, 255)
        cr.font.size = Pt(8.5)

    err_data = [
        ("Expected Return Date cannot be before Loan Date", "User selected return date earlier than loan start date.", "Set Expected Return Date to same day or later than Loan Effective Date."),
        ("Loan duration of X days exceeds maximum limit of 30 days", "Policy SLA breach: loan request exceeds 30-day corporate cap.", "Adjust return date within 30 days or obtain Department Head special waiver."),
        ("Asset AST-XXXX is currently 'Loaned Out' and unavailable", "Asset is currently borrowed by another staff member.", "Select a different asset marked Available in the inventory catalog."),
        ("Duplicate item AST-XXXX in loan table", "The same hardware serial was added twice in child table.", "Delete duplicate row using trash can icon."),
        ("Borrower EMP-XXXX has overdue equipment loan (LOAN-XXXX)", "Borrower has an unreturned loan past due date.", "STRICT LOCKOUT: Borrower must return overdue hardware before new loans can be issued."),
        ("At least one hardware item must be added to the loan", "User attempted to submit a loan with zero items.", "Click 'Add Row' and assign at least one hardware asset.")
    ]
    for idx, (msg, cause, fix) in enumerate(err_data):
        row = err_table.rows[idx + 1]
        bg = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate([msg, cause, fix]):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 35, 35, 50, 50)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(180, 40, 40)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 5: REST API
    add_sec_heading(doc, "5. Automated REST API Reference for Technicians", level=1)
    doc.add_paragraph(
        "For barcode scanning stations, automated locker check-ins, or IT ticketing systems:\n"
        "Endpoint: POST /api/method/test_project.api.loan_api.process_return\n"
        "Rate Limit: 30 requests / minute  |  Authentication: Token <api_key>:<api_secret>"
    ).paragraph_format.space_after = Pt(4)

    api_payload = (
        '{\n'
        '  "loan_id": "LOAN-2026-00042",\n'
        '  "actual_return_date": "2026-09-22",\n'
        '  "return_notes": "Returned via automated locker. All diagnostics passed."\n'
        '}'
    )
    add_code_block(doc, api_payload)

    doc.save(str(output_path))
    print(f"[SUCCESS] Built Working SOP Word document: {output_path}")

def convert_to_pdf(docx_path, pdf_path):
    ps_script = f"""
    $ErrorActionPreference = "Stop"
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $docx = "{str(docx_path.resolve())}"
    $pdf = "{str(pdf_path.resolve())}"
    try {{
        $doc = $word.Documents.Open($docx)
        $doc.SaveAs([ref]$pdf, [ref]17)
        $doc.Close()
        Write-Output "SUCCESS"
    }} finally {{
        $word.Quit()
        [System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null
    }}
    """
    
    ps_file = PROJECT_DOCS / "_convert_tmp.ps1"
    with open(ps_file, "w", encoding="utf-8") as f:
        f.write(ps_script)
        
    try:
        res = subprocess.run(
            ["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ps_file)],
            capture_output=True,
            text=True,
            check=True
        )
        print(f"[SUCCESS] Working SOP PDF generated: {pdf_path}")
    finally:
        if ps_file.exists():
            ps_file.unlink()

def main():
    docx_path = PROJECT_DOCS / "WORKING_SOP_EQUIPMENT_LOAN.docx"
    pdf_path = PROJECT_DOCS / "WORKING_SOP_EQUIPMENT_LOAN.pdf"
    
    build_docx(docx_path)
    convert_to_pdf(docx_path, pdf_path)
    
    user_downloads = Path.home() / "Downloads"
    if user_downloads.exists():
        import shutil
        shutil.copy2(docx_path, user_downloads / "WORKING_SOP_EQUIPMENT_LOAN.docx")
        shutil.copy2(pdf_path, user_downloads / "WORKING_SOP_EQUIPMENT_LOAN.pdf")
        print(f"[SUCCESS] Copied Working SOP to Downloads: {user_downloads / 'WORKING_SOP_EQUIPMENT_LOAN.pdf'}")

if __name__ == "__main__":
    main()
