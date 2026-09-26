#!/usr/bin/env python3
"""
Generates the comprehensive FRAPPE_ECC_SOP.docx (v1.1.0 with 20 Agents) and converts it to FRAPPE_ECC_SOP.pdf.
"""

import os
import sys
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = Path(__file__).resolve().parent.parent / "docs"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE", alert_type="note"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    border_colors = {"note": "2B579A", "warning": "D83B01", "tip": "107C41"}
    bg_colors = {"note": "F0F4F8", "warning": "FFF4CE", "tip": "EDF7ED"}
    
    fill_hex = bg_colors.get(alert_type, "F0F4F8")
    b_color = border_colors.get(alert_type, "2B579A")
    set_cell_background(cell, fill_hex)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{b_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor.from_string(b_color)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(code_text.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(36, 41, 46)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx(output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "STANDARD OPERATING PROCEDURE | FRAPPE ECC (20 AGENTS • 18 SKILLS • 20 COMMANDS)"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Calibri"
        hp.style.font.size = Pt(8.5)
        hp.style.font.color.rgb = RGBColor(128, 128, 128)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Document ID: SOP-ENG-FRAPPE-ECC-001  |  Version 1.1.0  |  Production Standard"
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.style.font.name = "Calibri"
        fp.style.font.size = Pt(8.5)
        fp.style.font.color.rgb = RGBColor(128, 128, 128)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("STANDARD OPERATING PROCEDURE (SOP)")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(31, 78, 121)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Engineering Coordination Center for Frappe Framework & ERPNext (Frappe ECC)")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(89, 89, 89)

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document ID", "SOP-ENG-FRAPPE-ECC-001"),
        ("Version & Release", "Version 1.1.0 (Full 20-Agent Engineering Release)"),
        ("Effective Date", "2026-09-18"),
        ("Applicability", "Frappe v14 / v15 / v16, ERPNext, Antigravity, Claude Code, Cursor, Codex")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        set_cell_background(c0, "E9EEF4")
        set_cell_margins(c0, 50, 50, 80, 80)
        set_cell_margins(c1, 50, 50, 80, 80)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_sop_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        return h

    add_sop_heading("1. Purpose & Scope", level=1)
    doc.add_paragraph(
        "This Standard Operating Procedure defines the mandatory engineering practices, execution phases, "
        "and automated quality controls for building, testing, and securing Frappe Framework applications using "
        "Frappe ECC. Frappe ECC equips your AI assistant with 20 specialist agents, 18 production skills, "
        "20 slash commands, and Frappe Shield."
    ).paragraph_format.space_after = Pt(6)

    add_sop_heading("2. System Architecture & Inventory", level=1)
    
    comp_table = doc.add_table(rows=6, cols=3)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Component Layer", "Count", "Core Responsibilities"]
    for col_idx, h in enumerate(headers):
        cell = comp_table.rows[0].cells[col_idx]
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 60, 60, 80, 80)
        cp = cell.paragraphs[0]
        cr = cp.add_run(h)
        cr.bold = True
        cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(255, 255, 255)
        cr.font.size = Pt(9.5)

    rows_data = [
        ("Specialist AI Agents", "20", "HLD, LLD, UI/UX, Wireframes, Clickable Prototypes, Fullstack Turnkey Dev, Manual QA, Automated Testing (Playwright), Planning, Controllers, Desk UI, TDD, Code Review, Security, DevOps, Migrations, Reporting, APIs, Docs"),
        ("Workflow Skills", "18", "HLD/LLD design, Wireframing & Prototyping, QA testing & Playwright automation, Turnkey scaffolding, DocType modeling, QueryBuilder, hooks.py, Desk scripts, REST API, Permissions, TDD, background jobs, reports, patches, Bench CLI, Frappe UI, Portal, Security Audit"),
        ("Slash Commands", "20", "/frappe:hld, /frappe:lld, /frappe:wireframe, /frappe:prototype, /frappe:build-e2e, /frappe:manual-qa, /frappe:e2e-test, /frappe:plan, /frappe:doctype, /frappe:controller, /frappe:client-script, /frappe:hook, /frappe:api, /frappe:test, /frappe:review, /frappe:security, /frappe:patch, /frappe:report, /frappe:bench, /frappe:help"),
        ("Coding Rules", "5", "Always-loaded standards: Core architecture, Python backend, Desk JS, Security/Permissions, and Database optimization"),
        ("Frappe Shield", "1 CLI", "AST static analyzer detecting SQL injection in frappe.db.sql, db.commit() violations, unvalidated guest APIs, and N+1 loops")
    ]
    for row_idx, r_data in enumerate(rows_data, 1):
        row = comp_table.rows[row_idx]
        bg = "F9FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate(r_data):
            c = row.cells[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 50, 50, 80, 80)
            cp = c.paragraphs[0]
            cr = cp.add_run(text)
            cr.font.name = "Calibri"
            cr.font.size = Pt(9)
            if c_idx == 0:
                cr.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_sop_heading("3. Installation Procedure", level=1)
    add_callout(doc, "Antigravity setup registers all 18 skills into your global config directory so every workspace automatically accesses Frappe domain intelligence.", "SETUP INSTRUCTION", "tip")
    
    doc.add_paragraph("Command for Linux / macOS / WSL:")
    add_code_block(doc, "./install.sh --profile full --target antigravity")
    
    doc.add_paragraph("Command for Windows PowerShell:")
    add_code_block(doc, ".\\install.ps1 -Profile full -Target antigravity")

    add_sop_heading("4. The 12-Phase Feature Development SOP", level=1)
    
    phases = [
        ("Phase 1: High-Level Design (HLD) & System Context",
         "Author system context, component topologies, and non-functional requirements (SLA, caching with Redis, multi-tenant bench topology) using '/frappe:hld'."),

        ("Phase 2: Low-Level Design (LLD) & Entity-Relationship Modeling",
         "Generate precise entity-relationship diagrams (Mermaid erDiagram), class hierarchies, and state transition tables using '/frappe:lld'."),

        ("Phase 3: Visual Wireframing & Clickable Interactive Prototyping",
         "Create visual ASCII/SVG wireframes using '/frappe:wireframe', followed by standalone clickable HTML/Vue prototypes using '/frappe:prototype' for stakeholder approval before code is written."),

        ("Phase 4: Schema Modeling & DocType Scaffolding",
         "Generate schema JSONs with sections, column breaks, indexes, and role permission arrays using '/frappe:doctype'."),

        ("Phase 5: Test-Driven Development (TDD) — Unit Tests First",
         "Draft failing test assertions in 'test_<doctype>.py' using FrappeTestCase. Confirm tests fail (RED) before implementing controllers."),

        ("Phase 6: Turnkey Fullstack Implementation (Zero Placeholders)",
         "Synthesize complete vertical implementations (controller lifecycle, Desk client scripts, dialogs, auto-math) using '/frappe:build-e2e'. Zero placeholders allowed."),

        ("Phase 7: Hook Registration & Event Orchestration",
         "Wire document events, background cron jobs, and class overrides into hooks.py using '/frappe:hook'."),

        ("Phase 8: Whitelisted APIs & Webhook Integrations",
         "Expose secure REST endpoints with role checks (frappe.only_for), rate limiting, and input sanitization using '/frappe:api'."),

        ("Phase 9: Database Schema Migrations & Patches",
         "Author idempotent migration patches in 'patches/' and register in 'patches.txt' using '/frappe:patch'."),

        ("Phase 10: Manual QA Matrix & Exploratory Verification",
         "Generate structured manual test scenarios, edge-case boundary tests, and release sign-off checklists using '/frappe:manual-qa'."),

        ("Phase 11: End-to-End Automated Browser Testing (Playwright)",
         "Author headless Playwright browser tests that click through Desk forms, fill fields, and assert submission badges using '/frappe:e2e-test'."),

        ("Phase 12: Security Audit with Frappe Shield & Code Review",
         "Execute AST static scanning with Frappe Shield ('/frappe:security') to ensure zero SQL injection, zero manual db.commit() calls, and zero security flaws.")
    ]

    for p_title, p_desc in phases:
        add_sop_heading(p_title, level=2)
        doc.add_paragraph(p_desc).paragraph_format.space_after = Pt(3)

    add_sop_heading("5. Master Command Cheat Sheet", level=1)
    
    cmd_table = doc.add_table(rows=21, cols=3)
    cmd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cmd_headers = ["Slash Command", "Specialist Agent", "Standard Operation"]
    for c_idx, h in enumerate(cmd_headers):
        cell = cmd_table.rows[0].cells[c_idx]
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, 60, 60, 80, 80)
        cp = cell.paragraphs[0]
        cr = cp.add_run(h)
        cr.bold = True
        cr.font.name = "Calibri"
        cr.font.color.rgb = RGBColor(255, 255, 255)
        cr.font.size = Pt(9.5)

    commands_summary = [
        ("/frappe:hld '<feature>'", "frappe-hld-architect", "Generate High-Level Design with C4 architecture diagrams"),
        ("/frappe:lld '<name>'", "frappe-lld-designer", "Generate Low-Level Design with ER & state machine diagrams"),
        ("/frappe:wireframe '<name>'", "frappe-wireframe-builder", "Generate visual ASCII / SVG wireframe layout mockups"),
        ("/frappe:prototype '<name>'", "frappe-interactive-prototyper", "Generate clickable interactive HTML/Vue prototype"),
        ("/frappe:plan '<feature>'", "frappe-planner", "Architect feature schema blueprint & module taxonomy"),
        ("/frappe:doctype '<Name>'", "frappe-planner", "Scaffold .json, .py, .js, and test_*.py files"),
        ("/frappe:build-e2e '<spec>'", "frappe-fullstack-developer", "Turnkey fullstack vertical synthesis (zero placeholders)"),
        ("/frappe:controller '<Name>'", "frappe-backend-builder", "Implement controller lifecycle hooks and validations"),
        ("/frappe:client-script '<Name>'", "frappe-desk-builder", "Build Desk form UI interactions & custom dialogs"),
        ("/frappe:hook [event|cron]", "frappe-architect", "Wire doc_events, scheduler cron, or class overrides"),
        ("/frappe:api '<name>'", "frappe-api-integrator", "Create secure whitelisted REST API endpoint"),
        ("/frappe:test '<Name>'", "frappe-tdd-guide", "Scaffold unit tests with FrappeTestCase"),
        ("/frappe:manual-qa '<Name>'", "frappe-manual-qa", "Generate manual test scenarios & QA sign-off checklist"),
        ("/frappe:e2e-test '<Name>'", "frappe-automated-tester", "Generate Playwright automated browser test script"),
        ("/frappe:review [path]", "frappe-code-reviewer", "Review diffs in fresh context for Frappe anti-patterns"),
        ("/frappe:security [path]", "frappe-security-reviewer", "Execute AST security scanner (SQLi, IDOR, XSS)"),
        ("/frappe:patch '<desc>'", "frappe-migration-patcher", "Draft idempotent database migration patch in patches.txt"),
        ("/frappe:report '<name>'", "frappe-report-builder", "Scaffold Script Report (Python backend + JS frontend)"),
        ("/frappe:bench [task]", "frappe-bench-devops", "Diagnose bench errors and site operations"),
        ("/frappe:help", "All Agents", "Display master quick cheat sheet")
    ]
    for r_idx, (cmd, agent, desc) in enumerate(commands_summary, 1):
        row = cmd_table.rows[r_idx]
        bg = "F9FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, text in enumerate([cmd, agent, desc]):
            c = row.cells[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 40, 40, 60, 60)
            cp = c.paragraphs[0]
            cr = cp.add_run(text)
            cr.font.name = "Calibri"
            cr.font.size = Pt(8.5)
            if c_idx == 0:
                cr.font.name = "Consolas"
                cr.bold = True
                cr.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    sign_table = doc.add_table(rows=1, cols=1)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_cell = sign_table.cell(0, 0)
    set_cell_background(s_cell, "EAECEE")
    set_cell_margins(s_cell, 80, 80, 120, 120)
    sp = s_cell.paragraphs[0]
    sr = sp.add_run("APPROVED & ENFORCED BY: Antigravity Systems Engineering Group\nDISTRIBUTION: Open Source / All Engineering Teams")
    sr.font.name = "Calibri"
    sr.font.size = Pt(8.5)
    sr.bold = True
    sr.font.color.rgb = RGBColor(80, 80, 80)

    doc.save(str(output_path))
    print(f"[SUCCESS] DOCX generated: {output_path}")

def convert_to_pdf(docx_path, pdf_path):
    ps_script = f"""
    $docx = "{str(docx_path).replace('\\', '\\\\')}"
    $pdf = "{str(pdf_path).replace('\\', '\\\\')}"
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    try {{
        $doc = $word.Documents.Open($docx)
        $doc.SaveAs([ref]$pdf, [ref]17)
        $doc.Close()
        Write-Output "SUCCESS"
    }} finally {{
        $word.Quit()
        [System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null
    }}
    """
    
    ps_file = DOCS_DIR / "_convert_tmp.ps1"
    with open(ps_file, "w", encoding="utf-8") as f:
        f.write(ps_script)
        
    try:
        res = subprocess.run(
            ["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ps_file)],
            capture_output=True,
            text=True,
            check=True
        )
        print(f"[SUCCESS] PDF generated: {pdf_path}")
    finally:
        if ps_file.exists():
            ps_file.unlink()

def main():
    docx_path = DOCS_DIR / "FRAPPE_ECC_SOP.docx"
    pdf_path = DOCS_DIR / "FRAPPE_ECC_SOP.pdf"
    
    build_docx(docx_path)
    convert_to_pdf(docx_path, pdf_path)
    
    user_downloads = Path.home() / "Downloads"
    if user_downloads.exists():
        import shutil
        shutil.copy2(docx_path, user_downloads / "FRAPPE_ECC_SOP.docx")
        shutil.copy2(pdf_path, user_downloads / "FRAPPE_ECC_SOP.pdf")
        print(f"[SUCCESS] Copied updated SOP to Downloads: {user_downloads / 'FRAPPE_ECC_SOP.pdf'}")

if __name__ == "__main__":
    main()
